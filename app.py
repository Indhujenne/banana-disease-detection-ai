from flask import Flask, render_template, request
import os
import numpy as np
from PIL import Image
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook
from huggingface_hub import hf_hub_download
from ai_edge_litert.interpreter import Interpreter


# ============================================================
# FLASK APP
# ============================================================

app = Flask(__name__)

UPLOAD_FOLDER = "static"
EXCEL_FILE = "results.xlsx"

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)


# ============================================================
# DOWNLOAD TFLITE MODEL FROM HUGGING FACE
# ============================================================

print("Downloading TFLite model from Hugging Face...")

MODEL_FILE = hf_hub_download(
    repo_id="jenneindhu/banana-disease-resnet50",
    filename="banana_disease_resnet50.tflite"
)

print("TFLite model downloaded successfully!")
print("Loading TFLite model...")


# ============================================================
# LOAD LITERT MODEL
# ============================================================

interpreter = Interpreter(model_path=MODEL_FILE)

interpreter.allocate_tensors()

input_details = interpreter.get_input_details()
output_details = interpreter.get_output_details()

print("TFLite model loaded successfully!")

print("Input shape:", input_details[0]["shape"])
print("Input type:", input_details[0]["dtype"])
print("Output shape:", output_details[0]["shape"])
print("Output type:", output_details[0]["dtype"])


# ============================================================
# CLASS NAMES
# IMPORTANT: KEEP THE SAME ORDER AS TRAINING
# ============================================================

CLASS_NAMES = [
    "black_sigatoka",
    "healthy",
    "moko_disease",
    "panama_disease",
    "yellow_sigatoka"
]


# ============================================================
# DISPLAY NAMES
# ============================================================

DISPLAY_NAMES = {
    "black_sigatoka": "Black Sigatoka Disease",
    "healthy": "Healthy Leaf",
    "moko_disease": "Moko Disease",
    "panama_disease": "Panama Disease",
    "yellow_sigatoka": "Yellow Sigatoka Disease"
}


# ============================================================
# ADVICE
# ============================================================

ADVICES = {

    "black_sigatoka": {
        "English": (
            "Remove infected leaves and maintain good field sanitation. "
            "Use recommended fungicides and avoid excessive moisture."
        ),
        "Telugu": (
            "వ్యాధి సోకిన ఆకులను తొలగించండి. పొలాన్ని శుభ్రంగా ఉంచండి. "
            "సిఫార్సు చేసిన శిలీంద్రనాశకాలను ఉపయోగించండి మరియు అధిక తేమను నివారించండి."
        ),
        "Hindi": (
            "संक्रमित पत्तियों को हटाएं और खेत को साफ रखें। "
            "अनुशंसित फफूंदनाशकों का उपयोग करें और अधिक नमी से बचें."
        ),
        "Tamil": (
            "பாதிக்கப்பட்ட இலைகளை அகற்றி வயலை சுத்தமாக வைத்திருக்கவும். "
            "பரிந்துரைக்கப்பட்ட பூஞ்சைக் கொல்லிகளை பயன்படுத்தவும்."
        )
    },

    "moko_disease": {
        "English": (
            "Remove and destroy infected plants. "
            "Maintain field hygiene and avoid spreading contaminated soil or tools."
        ),
        "Telugu": (
            "వ్యాధి సోకిన మొక్కలను తొలగించి నాశనం చేయండి. "
            "పొలంలో పరిశుభ్రత పాటించండి మరియు కలుషితమైన పనిముట్లను ఉపయోగించవద్దు."
        ),
        "Hindi": (
            "संक्रमित पौधों को हटाकर नष्ट करें। "
            "खेत की स्वच्छता बनाए रखें और दूषित उपकरणों से बचें."
        ),
        "Tamil": (
            "பாதிக்கப்பட்ட செடிகளை அகற்றி அழிக்கவும். "
            "வயல் சுகாதாரத்தை பராமரிக்கவும்."
        )
    },

    "panama_disease": {
        "English": (
            "Remove infected plants and avoid moving contaminated soil. "
            "Use healthy planting material and maintain proper drainage."
        ),
        "Telugu": (
            "వ్యాధి సోకిన మొక్కలను తొలగించండి. కలుషితమైన మట్టిని ఒక ప్రదేశం నుండి "
            "మరొక ప్రదేశానికి తరలించవద్దు. ఆరోగ్యకరమైన నాట్లను ఉపయోగించండి."
        ),
        "Hindi": (
            "संक्रमित पौधों को हटाएं और दूषित मिट्टी को न फैलाएं। "
            "स्वस्थ रोपण सामग्री का उपयोग करें और उचित जल निकासी रखें."
        ),
        "Tamil": (
            "பாதிக்கப்பட்ட செடிகளை அகற்றவும். "
            "சுத்தமான நடவு பொருட்களை பயன்படுத்தி நல்ல வடிகால் வசதி ஏற்படுத்தவும்."
        )
    },

    "yellow_sigatoka": {
        "English": (
            "Remove severely infected leaves and maintain good air circulation. "
            "Use recommended fungicide treatment when necessary."
        ),
        "Telugu": (
            "తీవ్రంగా వ్యాధి సోకిన ఆకులను తొలగించండి. "
            "మొక్కల మధ్య మంచి గాలి ప్రసరణ ఉండేలా చూసుకోండి."
        ),
        "Hindi": (
            "बहुत अधिक संक्रमित पत्तियों को हटाएं और हवा का अच्छा संचार बनाए रखें। "
            "आवश्यक होने पर अनुशंसित फफूंदनाशक का उपयोग करें."
        ),
        "Tamil": (
            "கடுமையாக பாதிக்கப்பட்ட இலைகளை அகற்றி நல்ல காற்றோட்டத்தை உறுதி செய்யவும்."
        )
    }
}


# ============================================================
# FERTILIZER RECOMMENDATIONS
# ============================================================

FERTILIZERS = {

    "black_sigatoka": {
        "English": "Use balanced NPK fertilizer and maintain adequate potassium.",
        "Telugu": "சమతుల్యమైన NPK ఎరువును ఉపయోగించి తగినంత పొటాషియం అందించండి.",
        "Hindi": "संतुलित NPK उर्वरक का उपयोग करें और पर्याप्त पोटैशियम दें.",
        "Tamil": "சமநிலை NPK உரத்தை பயன்படுத்தி போதுமான பொட்டாசியம் வழங்கவும்."
    },

    "moko_disease": {
        "English": "Use balanced fertilizer based on soil requirements and maintain plant nutrition.",
        "Telugu": "మట్టి అవసరాలకు అనుగుణంగా సమతుల్య ఎరువును ఉపయోగించండి.",
        "Hindi": "मिट्टी की आवश्यकता के अनुसार संतुलित उर्वरक का उपयोग करें.",
        "Tamil": "மண் தேவைக்கேற்ப சமநிலை உரத்தை பயன்படுத்தவும்."
    },

    "panama_disease": {
        "English": "Maintain balanced nutrition with adequate potassium and organic matter.",
        "Telugu": "తగినంత పొటాషియం మరియు సేంద్రీయ పదార్థంతో సమతుల్య పోషకాలను అందించండి.",
        "Hindi": "पर्याप्त पोटैशियम और जैविक पदार्थ के साथ संतुलित पोषण दें.",
        "Tamil": "போதுமான பொட்டாசியம் மற்றும் இயற்கை பொருட்களுடன் சமநிலை ஊட்டச்சத்தை வழங்கவும்."
    },

    "yellow_sigatoka": {
        "English": "Use balanced NPK fertilizer and maintain adequate micronutrients.",
        "Telugu": "సమతుల్య NPK ఎరువును ఉపయోగించి అవసరమైన సూక్ష్మ పోషకాలను అందించండి.",
        "Hindi": "संतुलित NPK उर्वरक और आवश्यक सूक्ष्म पोषक तत्व दें.",
        "Tamil": "சமநிலை NPK உரம் மற்றும் தேவையான நுண்ணூட்டச்சத்துக்களை வழங்கவும்."
    }
}


# ============================================================
# HEALTHY LEAF ADVICE
# ============================================================

HEALTHY_ADVICE = {
    "English": (
        "The banana leaf appears healthy. Continue proper irrigation, "
        "balanced fertilization and regular monitoring."
    ),
    "Telugu": (
        "అరటి ఆకు ఆరోగ్యంగా కనిపిస్తోంది. సరైన నీటి పారుదల, "
        "సమతుల్య ఎరువులు మరియు క్రమం తప్పకుండా పర్యవేక్షణ కొనసాగించండి."
    ),
    "Hindi": (
        "केले का पत्ता स्वस्थ दिखाई दे रहा है। "
        "उचित सिंचाई, संतुलित उर्वरक और नियमित निगरानी जारी रखें."
    ),
    "Tamil": (
        "வாழை இலை ஆரோக்கியமாக தெரிகிறது. "
        "சரியான நீர்ப்பாசனம் மற்றும் சமநிலை உரமிடலை தொடரவும்."
    )
}


# ============================================================
# CREATE EXCEL FILE IF IT DOES NOT EXIST
# ============================================================

if not os.path.exists(EXCEL_FILE):

    workbook = Workbook()

    sheet = workbook.active

    sheet.append([
        "Image",
        "Disease",
        "Confidence",
        "Language",
        "Advice",
        "Fertilizer"
    ])

    workbook.save(EXCEL_FILE)


# ============================================================
# HOME PAGE
# ============================================================

@app.route("/")
def home():

    previous = []

    try:

        workbook = load_workbook(EXCEL_FILE)

        sheet = workbook.active

        for row in sheet.iter_rows(min_row=2, values_only=True):

            if row[0]:

                previous.append({
                    "image": row[0],
                    "disease": row[1],
                    "confidence": row[2],
                    "language": row[3],
                    "advice": row[4],
                    "fertilizer": row[5]
                })

        workbook.close()

    except Exception as error:

        print("Excel read error:", error)

    return render_template(
        "index.html",
        previous=previous
    )


# ============================================================
# PREDICTION FUNCTION
# ============================================================

def predict_image(image):

    # Convert image to RGB
    image = image.convert("RGB")

    # Resize to model input size
    image = image.resize((224, 224))

    # Convert to numpy
    image_array = np.array(
        image,
        dtype=np.float32
    )

    # Add batch dimension
    image_array = np.expand_dims(
        image_array,
        axis=0
    )

    # Check expected input type
    expected_dtype = input_details[0]["dtype"]

    if image_array.dtype != expected_dtype:

        image_array = image_array.astype(
            expected_dtype
        )

    # Set input tensor
    interpreter.set_tensor(
        input_details[0]["index"],
        image_array
    )

    # Run model
    interpreter.invoke()

    # Get prediction
    predictions = interpreter.get_tensor(
        output_details[0]["index"]
    )

    # Get highest probability
    predicted_index = int(
        np.argmax(predictions[0])
    )

    confidence = float(
        predictions[0][predicted_index]
    ) * 100

    predicted_class = CLASS_NAMES[predicted_index]

    return predicted_class, confidence


# ============================================================
# PREDICT ROUTE
# ============================================================

@app.route("/predict", methods=["POST"])
def predict():

    try:

        files = request.files.getlist("images")

        language = request.form.get(
            "language",
            "English"
        )

        if not files:

            return {
                "error": "No image uploaded"
            }, 400

        results = []

        workbook = load_workbook(EXCEL_FILE)

        sheet = workbook.active

        for file in files:

            if not file or file.filename == "":
                continue

            filename = secure_filename(
                file.filename
            )

            filepath = os.path.join(
                app.config["UPLOAD_FOLDER"],
                filename
            )

            file.save(filepath)

            # Open uploaded image
            image = Image.open(filepath)

            # Predict
            predicted_class, confidence = predict_image(
                image
            )

            # Display name
            disease_name = DISPLAY_NAMES.get(
                predicted_class,
                predicted_class
            )

            # Advice
            if predicted_class == "healthy":

                advice = HEALTHY_ADVICE.get(
                    language,
                    HEALTHY_ADVICE["English"]
                )

                fertilizer = "No special fertilizer required."

            else:

                advice = ADVICES.get(
                    predicted_class,
                    {}
                ).get(
                    language,
                    ADVICES.get(
                        predicted_class,
                        {}
                    ).get(
                        "English",
                        "No advice available."
                    )
                )

                fertilizer = FERTILIZERS.get(
                    predicted_class,
                    {}
                ).get(
                    language,
                    FERTILIZERS.get(
                        predicted_class,
                        {}
                    ).get(
                        "English",
                        "No fertilizer recommendation available."
                    )
                )

            # Save result to Excel
            sheet.append([
                filename,
                disease_name,
                round(confidence, 2),
                language,
                advice,
                fertilizer
            ])

            results.append({
                "image": filename,
                "disease": disease_name,
                "confidence": round(
                    confidence,
                    2
                ),
                "language": language,
                "advice": advice,
                "fertilizer": fertilizer
            })

        workbook.save(EXCEL_FILE)

        workbook.close()

        return {
            "results": results
        }

    except Exception as error:

        print("Prediction error:", error)

        return {
            "error": str(error)
        }, 500


# ============================================================
# RUN APPLICATION
# ============================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=int(
            os.environ.get(
                "PORT",
                5000
            )
        ),
        debug=False
    )